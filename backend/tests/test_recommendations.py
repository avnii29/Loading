from decimal import Decimal
from pathlib import Path

import openpyxl

from app.database import SessionLocal
from app.models.organization import DonationCentre
from app.services.centre_import import import_workbook
from app.services.recommendations import recommend_centres


def test_radius_keeps_a_destination_10km_away_and_skips_other_categories():
    db = SessionLocal()
    near = DonationCentre(
        name="Test Near Kitchen",
        listing="published",
        is_active=True,
        latitude=Decimal("17.450000"),
        longitude=Decimal("78.450000"),
        approximate=True,
        accepted_categories=["FOOD"],
        geocode_status="existing_approximate",
        operational_status="location_needs_confirmation",
        source_key="published:test-near-kitchen",
    )
    far = DonationCentre(
        name="Test Far Books",
        listing="published",
        is_active=True,
        latitude=Decimal("17.900000"),
        longitude=Decimal("79.200000"),
        approximate=True,
        accepted_categories=["BOOKS"],
        geocode_status="existing_approximate",
        operational_status="location_needs_confirmation",
        source_key="published:test-far-books",
    )
    db.add_all([near, far])
    db.commit()
    try:
        items, search = recommend_centres(db, latitude=17.385, longitude=78.4867, category="FOOD", limit=30)
        near = next(item for item in items if item.centre.name == "Test Near Kitchen")
        assert search != "none"
        assert near.distance_km < 12
        assert all(item.centre.name != "Test Far Books" for item in items)
    finally:
        db.query(DonationCentre).filter(DonationCentre.source_key.like("published:test-%")).delete(synchronize_session=False)
        db.commit()
        db.close()


def test_import_keeps_leads_and_removed_inactive(tmp_path: Path):
    book = tmp_path / "centres.xlsx"
    wb = openpyxl.Workbook()
    main = wb.active
    main.title = "Facilities (verified sites)"
    main.append(["S.No", "NGO / Facility", "Facility Type", "District / Area", "Address (as published)", "Phone", "What it provides", "People served / capacity", "What to donate (suggested; confirm first)", "Source", "Data caveat"])
    main.append([1, "Import Sample Home", "Shelter", "Hyderabad", "Hyderabad", None, "Food", "10", "Food", "https://example.test", "Confirm first"])
    leads = wb.create_sheet("Unverified leads")
    leads.append(["S.No", "Name", "Location / Address found", "Why it is not in the main list", "Source"])
    leads.append([1, "Import Lead", "Somewhere", "Not confirmed", "https://example.test/lead"])
    removed = wb.create_sheet("Removed (HQ-only)")
    removed.append(["S.No", "NGO from earlier sheet", "Reason removed"])
    removed.append([1, "Import HQ", "Headquarters only"])
    wb.save(book)

    db = SessionLocal()
    try:
        first = import_workbook(db, book)
        db.commit()
        second = import_workbook(db, book)
        db.commit()
        assert first.imported == 3
        assert second.imported == 0
        assert second.updated == 3
        assert first.leads_kept_separate == 1
        assert first.removed_not_activated == 1
        lead = db.query(DonationCentre).filter(DonationCentre.name == "Import Lead").one()
        gone = db.query(DonationCentre).filter(DonationCentre.name == "Import HQ").one()
        assert lead.is_active is False
        assert gone.is_active is False
        assert gone.listing == "removed"
        home = db.query(DonationCentre).filter(DonationCentre.name == "Import Sample Home").one()
        assert home.phone is None
        assert home.latitude is None
        assert home.operational_status == "research_imported"
    finally:
        db.query(DonationCentre).filter(DonationCentre.name.in_(["Import Sample Home", "Import Lead", "Import HQ"])).delete(synchronize_session=False)
        db.commit()
        db.close()
