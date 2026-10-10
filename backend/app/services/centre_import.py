"""Idempotent import of the facility workbook.

The main sheet is research, not Anuvartan operational verification.
This importer does not call a geocoder and does not invent coordinates or phone numbers.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.organization import DonationCentre

ACTIVE_SHEET = "Facilities (verified sites)"
LEAD_SHEET = "Unverified leads"
REMOVED_SHEET = "Removed (HQ-only)"


@dataclass
class ImportReport:
    imported: int = 0
    updated: int = 0
    skipped: int = 0
    duplicates: int = 0
    unresolved: int = 0
    invalid: int = 0
    leads_kept_separate: int = 0
    removed_not_activated: int = 0
    notes: list[str] = field(default_factory=list)


def _text(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).replace("\xa0", " ").strip()
    return text or None


def _key(listing: str, name: str) -> str:
    return f"{listing}:{re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-')}"[:320]


def suggested_categories(text: str | None) -> list[str]:
    value = (text or "").lower()
    found = []
    if any(word in value for word in ("food", "meal", "ration", "rice")):
        found.append("FOOD")
    if any(word in value for word in ("cloth", "uniform")):
        found.append("CLOTHES")
    if any(word in value for word in ("book", "school", "librar")):
        found.append("BOOKS")
    if "toy" in value:
        found.append("TOYS")
    if any(word in value for word in ("misc", "other", "household")):
        found.append("OTHER")
    return found


def _ambiguous(address: str | None, district: str | None) -> bool:
    if not address:
        return True
    compact = address.lower().strip(" ,")
    area = (district or "").lower().strip(" ,")
    if len(compact) < 12:
        return True
    return bool(area) and compact == area


def import_workbook(db: Session, path: Path) -> ImportReport:
    report = ImportReport()
    if not path.exists():
        report.invalid += 1
        report.notes.append(f"Workbook not found: {path}")
        return report
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    missing = [name for name in (ACTIVE_SHEET, LEAD_SHEET, REMOVED_SHEET) if name not in wb.sheetnames]
    if missing:
        report.invalid += 1
        report.notes.append(f"Missing sheets: {', '.join(missing)}")
        return report

    seen: set[str] = set()

    def upsert(sheet: str, listing: str, active: bool) -> None:
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            report.invalid += 1
            return
        for offset, row in enumerate(rows, start=2):
            if not row or not any(row):
                report.skipped += 1
                continue
            name = _text(row[1] if len(row) > 1 else None)
            if not name:
                report.invalid += 1
                continue
            key = _key(listing, name)
            if key in seen:
                report.duplicates += 1
                continue
            seen.add(key)
            address = _text(row[4] if listing != "removed" and len(row) > 4 else (row[2] if len(row) > 2 else None))
            if listing == "removed":
                address = _text(row[2] if len(row) > 2 else None)
            district = _text(row[3] if listing == "published" and len(row) > 3 else None)
            if listing == "unconfirmed":
                address = _text(row[2] if len(row) > 2 else None)
                district = None
            existing = db.scalar(select(DonationCentre).where(DonationCentre.source_key == key))
            if existing is None:
                existing = db.scalar(
                    select(DonationCentre).where(DonationCentre.name == name, DonationCentre.listing == listing)
                )
            ambiguous = _ambiguous(address, district)
            if existing is None:
                centre = DonationCentre(
                    name=name[:240],
                    listing=listing,
                    is_active=active,
                    source_key=key,
                    approximate=True,
                    geocode_status="needs_review" if ambiguous else "unresolved",
                    operational_status="not_active" if not active else "research_imported",
                )
                db.add(centre)
                report.imported += 1
                if centre.geocode_status != "resolved":
                    report.unresolved += 1
            else:
                centre = existing
                report.updated += 1
                if centre.latitude is None:
                    report.unresolved += 1
            centre.source_sheet = sheet
            centre.source_row = offset
            centre.source_key = key
            centre.listing = listing
            centre.is_active = active
            centre.name = name[:240]
            if listing == "published":
                centre.facility_type = _text(row[2] if len(row) > 2 else None)
                centre.district = district
                centre.address = address[:400] if address else None
                centre.phone = _text(row[5] if len(row) > 5 else None)
                centre.provides = _text(row[6] if len(row) > 6 else None)
                centre.capacity_notes = _text(row[7] if len(row) > 7 else None)
                centre.donate_notes = _text(row[8] if len(row) > 8 else None)
                centre.source_url = _text(row[9] if len(row) > 9 else None)
                centre.caveat = _text(row[10] if len(row) > 10 else None)
                centre.accepted_categories = suggested_categories(centre.donate_notes)
            elif listing == "unconfirmed":
                centre.address = address[:400] if address else None
                centre.caveat = _text(row[3] if len(row) > 3 else None)
                centre.source_url = _text(row[4] if len(row) > 4 else None)
            else:
                centre.caveat = address
                centre.operational_status = "not_active"
            if centre.latitude is not None and centre.longitude is not None:
                centre.geocode_status = "existing_approximate" if centre.approximate else "resolved"
                if active and centre.operational_status == "research_imported":
                    centre.operational_status = "location_needs_confirmation" if centre.approximate else "location_resolved"
            elif active:
                centre.geocode_status = "needs_review" if ambiguous else "unresolved"
                centre.operational_status = "research_imported"
            db.add(centre)

    upsert(ACTIVE_SHEET, "published", True)
    before_leads = report.imported + report.updated
    upsert(LEAD_SHEET, "unconfirmed", False)
    report.leads_kept_separate = report.imported + report.updated - before_leads
    before_removed = report.imported + report.updated
    upsert(REMOVED_SHEET, "removed", False)
    report.removed_not_activated = report.imported + report.updated - before_removed
    db.flush()
    return report
